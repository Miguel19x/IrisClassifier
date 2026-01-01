"""
Master Table API Router.

Provides endpoints for the unified Master Table (Gestión Listados).
Supports dual view modes (Enterprise/Client) with editable fields.
"""
import logging
from typing import Optional, List, Literal
from decimal import Decimal
from fastapi import APIRouter, Depends, Query, HTTPException, Body
from sqlalchemy.orm import Session
from sqlalchemy import func, desc, asc
import base64
import io

from database.connection import get_db
from database.models import MasterProduct, PriceList
from api.dependencies import get_current_user_id
from pydantic import BaseModel, Field

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/master-products", tags=["master-products"])


# Schemas
class MasterProductResponse(BaseModel):
    id: int
    index_number: int
    clean_code: str
    description: str
    brand: Optional[str]
    price_usd: Decimal
    review_status: str
    confidence_score: float
    source_list_id: int
    original_list_name: str
    margin_percentage: Optional[Decimal]
    final_price: Optional[Decimal]
    
    class Config:
        from_attributes = True


class MasterProductListResponse(BaseModel):
    products: List[MasterProductResponse]
    total: int
    page: int
    limit: int
    has_next: bool
    has_prev: bool


class UpdateMarginRequest(BaseModel):
    margin_percentage: Decimal = Field(..., ge=-100, le=1000)


class UpdateFinalPriceRequest(BaseModel):
    final_price: Decimal = Field(..., gt=0)


class PriceStatsResponse(BaseModel):
    """Statistics for heatmap percentile calculation."""
    min_price: Decimal
    max_price: Decimal
    p25: Decimal  # 25th percentile
    p50: Decimal  # 50th percentile (median)
    p75: Decimal  # 75th percentile


# Export Configuration Schemas
class ManualHeaderConfig(BaseModel):
    """Manual header configuration with rich text."""
    content: str  # HTML content from rich text editor
    alignment: Literal['left', 'center', 'right'] = 'center'


class TemplateFileConfig(BaseModel):
    """Template file configuration."""
    name: str
    type: Literal['pdf', 'excel']
    data: str  # Base64 encoded file content


class ExportConfigRequest(BaseModel):
    """Full export configuration for advanced exports."""
    filename: str = 'Listado_de_Productos'
    header_mode: Literal['manual', 'template'] = 'manual'
    manual_header: Optional[ManualHeaderConfig] = None
    template_file: Optional[TemplateFileConfig] = None
    view_mode: Literal['enterprise', 'client'] = 'enterprise'
    sort_by: Literal['index', 'alphabetical', 'brand', 'description', 'price'] = 'alphabetical'
    brand_filter: Optional[str] = None
    review_status_filter: Optional[Literal['pending', 'confirmed', 'rejected']] = None
    format: Literal['excel', 'pdf'] = 'excel'


# Endpoints
@router.get("", response_model=MasterProductListResponse)
async def get_master_products(
    view_mode: str = Query("enterprise", regex="^(enterprise|client)$"),
    sort_by: str = Query("alphabetical", regex="^(index|alphabetical|brand|description|price)$"),
    search: Optional[str] = Query(None, min_length=1, max_length=100),
    brand_filter: Optional[str] = None,
    review_status_filter: Optional[str] = Query(None, regex="^(pending|confirmed|rejected)$"),
    page: int = Query(1, ge=1),
    limit: int = Query(200, ge=1, le=500),
    user_id: int = Depends(get_current_user_id),
    db: Session = Depends(get_db)
):
    """
    Get master products with filters and pagination.
    
    Supports:
    - search: Text search across code, description, and brand
    - sort_by: Sort order (alphabetical default)
    - brand_filter: Filter by brand name
    - review_status_filter: Filter by review status
    """
    # Base query - filter by user's lists
    query = db.query(MasterProduct).join(
        PriceList, MasterProduct.source_list_id == PriceList.id
    ).filter(PriceList.user_id == user_id)
    
    # Apply search filter (searches code, description, and brand)
    if search:
        search_term = f"%{search.lower()}%"
        query = query.filter(
            (func.lower(MasterProduct.clean_code).like(search_term)) |
            (func.lower(MasterProduct.description).like(search_term)) |
            (func.lower(MasterProduct.brand).like(search_term))
        )
    
    # Apply brand filter
    if brand_filter:
        query = query.filter(
            func.lower(MasterProduct.brand).like(f"%{brand_filter.lower()}%")
        )
    
    if review_status_filter:
        query = query.filter(MasterProduct.review_status == review_status_filter)
    
    # Apply sorting (with secondary sort by index_number for consistency)
    if sort_by == "index":
        query = query.order_by(asc(MasterProduct.index_number))
    elif sort_by == "alphabetical":
        query = query.order_by(asc(MasterProduct.description), asc(MasterProduct.index_number))
    elif sort_by == "brand":
        query = query.order_by(asc(MasterProduct.brand), asc(MasterProduct.description), asc(MasterProduct.index_number))
    elif sort_by == "description":
        query = query.order_by(asc(MasterProduct.description), asc(MasterProduct.index_number))
    elif sort_by == "price":
        query = query.order_by(asc(MasterProduct.price_usd), asc(MasterProduct.index_number))
    
    # Get total count
    total = query.count()
    
    # Pagination
    offset = (page - 1) * limit
    products = query.offset(offset).limit(limit).all()
    
    return MasterProductListResponse(
        products=products,
        total=total,
        page=page,
        limit=limit,
        has_next=offset + limit < total,
        has_prev=page > 1
    )


@router.get("/stats", response_model=PriceStatsResponse)
async def get_price_stats(
    user_id: int = Depends(get_current_user_id),
    db: Session = Depends(get_db)
):
    """
    Get price statistics for heatmap percentile calculation.
    
    Returns min, max, and percentiles (p25, p50, p75) for dynamic heatmap coloring.
    """
    # Get all prices for user's products
    prices_query = db.query(MasterProduct.price_usd).join(
        PriceList, MasterProduct.source_list_id == PriceList.id
    ).filter(PriceList.user_id == user_id).order_by(MasterProduct.price_usd)
    
    prices = [p[0] for p in prices_query.all()]
    
    if not prices:
        return PriceStatsResponse(
            min_price=Decimal("0"),
            max_price=Decimal("0"),
            p25=Decimal("0"),
            p50=Decimal("0"),
            p75=Decimal("0")
        )
    
    # Calculate percentiles
    def percentile(data: List[Decimal], p: float) -> Decimal:
        k = (len(data) - 1) * p
        f = int(k)
        c = k - f
        if f + 1 < len(data):
            return data[f] * Decimal(1 - c) + data[f + 1] * Decimal(c)
        return data[f]
    
    return PriceStatsResponse(
        min_price=prices[0],
        max_price=prices[-1],
        p25=percentile(prices, 0.25),
        p50=percentile(prices, 0.50),
        p75=percentile(prices, 0.75)
    )


@router.patch("/{product_id}/margin")
async def update_margin(
    product_id: int,
    request: UpdateMarginRequest,
    user_id: int = Depends(get_current_user_id),
    db: Session = Depends(get_db)
):
    """
    Update margin percentage (Enterprise View).
    
    Automatically recalculates final_price = price_usd * (1 + margin/100)
    """
    product = db.query(MasterProduct).join(
        PriceList, MasterProduct.source_list_id == PriceList.id
    ).filter(
        MasterProduct.id == product_id,
        PriceList.user_id == user_id
    ).first()
    
    if not product:
        raise HTTPException(status_code=404, detail="Product not found")
    
    # Update margin
    product.margin_percentage = request.margin_percentage
    
    # Calculate final price
    margin_multiplier = 1 + (request.margin_percentage / 100)
    product.final_price = product.price_usd * margin_multiplier
    
    db.commit()
    db.refresh(product)
    
    logger.info(
        "margin_updated",
        extra={
            "product_id": product_id,
            "margin": float(request.margin_percentage),
            "final_price": float(product.final_price)
        }
    )
    
    return {"status": "success", "final_price": product.final_price}


@router.patch("/{product_id}/final-price")
async def update_final_price(
    product_id: int,
    request: UpdateFinalPriceRequest,
    user_id: int = Depends(get_current_user_id),
    db: Session = Depends(get_db)
):
    """
    Update final price (Client View).
    
    Automatically recalculates margin % = ((final - base) / base) * 100
    """
    product = db.query(MasterProduct).join(
        PriceList, MasterProduct.source_list_id == PriceList.id
    ).filter(
        MasterProduct.id == product_id,
        PriceList.user_id == user_id
    ).first()
    
    if not product:
        raise HTTPException(status_code=404, detail="Product not found")
    
    # Update final price
    product.final_price = request.final_price
    
    # Calculate margin percentage (inverse)
    if product.price_usd > 0:
        margin = ((request.final_price - product.price_usd) / product.price_usd) * 100
        product.margin_percentage = margin
    else:
        product.margin_percentage = Decimal("0")
    
    db.commit()
    db.refresh(product)
    
    logger.info(
        "final_price_updated",
        extra={
            "product_id": product_id,
            "final_price": float(request.final_price),
            "margin": float(product.margin_percentage)
        }
    )
    
    return {"status": "success", "margin_percentage": product.margin_percentage}


@router.patch("/{product_id}/review-status")
async def update_review_status(
    product_id: int,
    status: str = Query(..., regex="^(confirmed|rejected)$"),
    user_id: int = Depends(get_current_user_id),
    db: Session = Depends(get_db)
):
    """
    Update review status for a product.
    
    Used when user manually confirms or rejects a product with low confidence.
    """
    product = db.query(MasterProduct).join(
        PriceList, MasterProduct.source_list_id == PriceList.id
    ).filter(
        MasterProduct.id == product_id,
        PriceList.user_id == user_id
    ).first()
    
    if not product:
        raise HTTPException(status_code=404, detail="Product not found")
    
    product.review_status = status
    
    # If confirmed, update code registry
    if status == "confirmed":
        from services.etl_intelligent_service import ETLIntelligentService
        etl = ETLIntelligentService(db)
        await etl._update_code_registry(
            product.clean_code,
            product.description,
            product.brand
        )
    
    db.commit()
    
    logger.info(
        "review_status_updated",
        extra={
            "product_id": product_id,
            "status": status
        }
    )
    
    return {"status": "success"}


@router.get("/export")
async def export_master_products(
    format: str = Query("excel", regex="^(excel|pdf)$"),
    view_mode: str = Query("enterprise", regex="^(enterprise|client)$"),
    sort_by: str = Query("alphabetical", regex="^(index|alphabetical|brand|description|price)$"),
    custom_title: str = Query("Listado de Productos", description="Custom title for the export"),
    brand_filter: Optional[str] = None,
    review_status_filter: Optional[str] = Query(None, regex="^(pending|confirmed|rejected)$"),
    user_id: int = Depends(get_current_user_id),
    db: Session = Depends(get_db)
):
    """
    Export master products to Excel or PDF.
    
    Respects current filters and view mode:
    - enterprise: Includes margin % column
    - client: Includes final price column only
    
    Returns file download.
    """
    from fastapi.responses import StreamingResponse
    import io
    
    try:
        # Build query with filters (same as get_master_products)
        query = db.query(MasterProduct).join(
            PriceList, MasterProduct.source_list_id == PriceList.id
        ).filter(PriceList.user_id == user_id)
        
        if brand_filter:
            query = query.filter(
                func.lower(MasterProduct.brand).like(f"%{brand_filter.lower()}%")
            )
        
        if review_status_filter:
            query = query.filter(MasterProduct.review_status == review_status_filter)
        
        # Apply sorting (with secondary sort by index_number for consistency)
        if sort_by == "index":
            query = query.order_by(asc(MasterProduct.index_number))
        elif sort_by == "alphabetical":
            query = query.order_by(asc(MasterProduct.description), asc(MasterProduct.index_number))
        elif sort_by == "brand":
            query = query.order_by(asc(MasterProduct.brand), asc(MasterProduct.description), asc(MasterProduct.index_number))
        elif sort_by == "description":
            query = query.order_by(asc(MasterProduct.description), asc(MasterProduct.index_number))
        elif sort_by == "price":
            query = query.order_by(asc(MasterProduct.price_usd), asc(MasterProduct.index_number))
        
        products = query.all()
        logger.info(f"Exporting {len(products)} products as {format} in {view_mode} mode")
        
        if format == "excel":
            return await _export_to_excel(products, view_mode, custom_title)
        else:
            return await _export_to_pdf(products, view_mode, custom_title)
    except Exception as e:
        import traceback
        print(f"[EXPORT ERROR] {e}")
        traceback.print_exc()
        logger.error(f"Export failed: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=f"Export failed: {str(e)}")


@router.post("/export-advanced")
async def export_master_products_advanced(
    config: ExportConfigRequest = Body(...),
    user_id: int = Depends(get_current_user_id),
    db: Session = Depends(get_db)
):
    """
    Advanced export with template support.
    
    Supports two modes:
    - manual: Uses rich text header from config
    - template: Uses uploaded PDF/Excel as header/base
    
    Returns file download.
    """
    from fastapi.responses import StreamingResponse
    
    try:
        # Build query with filters
        query = db.query(MasterProduct).join(
            PriceList, MasterProduct.source_list_id == PriceList.id
        ).filter(PriceList.user_id == user_id)
        
        if config.brand_filter:
            query = query.filter(
                func.lower(MasterProduct.brand).like(f"%{config.brand_filter.lower()}%")
            )
        
        if config.review_status_filter:
            query = query.filter(MasterProduct.review_status == config.review_status_filter)
        
        # Apply sorting
        sort_map = {
            "index": [asc(MasterProduct.index_number)],
            "alphabetical": [asc(MasterProduct.description), asc(MasterProduct.index_number)],
            "brand": [asc(MasterProduct.brand), asc(MasterProduct.description), asc(MasterProduct.index_number)],
            "description": [asc(MasterProduct.description), asc(MasterProduct.index_number)],
            "price": [asc(MasterProduct.price_usd), asc(MasterProduct.index_number)],
        }
        for order in sort_map.get(config.sort_by, sort_map["alphabetical"]):
            query = query.order_by(order)
        
        products = query.all()
        logger.info(f"Advanced export: {len(products)} products as {config.format} in {config.view_mode} mode")
        
        # Route based on template or manual header
        if config.header_mode == "template" and config.template_file:
            # Template-based export
            if config.format == "excel":
                return await _export_to_excel_with_template(
                    products, config.view_mode, config.template_file, config.filename
                )
            else:
                return await _export_to_pdf_with_template(
                    products, config.view_mode, config.template_file, config.filename
                )
        else:
            # Manual header export - pass HTML content and filename separately
            header_html = config.manual_header.content if config.manual_header else ""
            print(f"[DEBUG] Header HTML received: {repr(header_html)}")
            if config.format == "excel":
                return await _export_to_excel_with_header(
                    products, config.view_mode, header_html, config.filename
                )
            else:
                return await _export_to_pdf_with_header(
                    products, config.view_mode, header_html, config.filename
                )
                
    except Exception as e:
        import traceback
        print(f"[ADVANCED EXPORT ERROR] {e}")
        traceback.print_exc()
        logger.error(f"Advanced export failed: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=f"Export failed: {str(e)}")


async def _export_to_excel(products: List[MasterProduct], view_mode: str, custom_title: str = "Listado de Productos"):
    """Generate Excel export."""
    from fastapi.responses import StreamingResponse
    import io
    import re
    
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise HTTPException(
            status_code=500, 
            detail="openpyxl not installed. Run: pip install openpyxl"
        )
    
    wb = openpyxl.Workbook()
    ws = wb.active
    
    # Strip HTML tags for sheet name
    sheet_name = re.sub(r'<[^>]+>', '', custom_title).strip()[:31]
    ws.title = sheet_name if sheet_name else "Listado"
    
    # Parse HTML formatting from title
    is_bold = '<strong>' in custom_title or '<b>' in custom_title
    is_italic = '<em>' in custom_title or '<i>' in custom_title
    is_underline = '<u>' in custom_title
    is_centered = 'text-align: center' in custom_title or 'text-align:center' in custom_title
    
    # Clean title text (strip HTML)
    title_text = re.sub(r'<[^>]+>', '', custom_title).strip()
    
    # Add title row (row 1) - merged across all columns
    if view_mode == "enterprise":
        headers = ["N°", "CÓDIGO O REFERENCIA", "DESCRIPCIÓN", "MARCA", "LISTA ORIGINAL", "USD", "%", "USD FINAL"]
    else:
        headers = ["N°", "CÓDIGO O REFERENCIA", "DESCRIPCIÓN", "MARCA", "USD"]
    
    # Merge cells for title (row 1)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = ws.cell(row=1, column=1, value=title_text)
    title_cell.font = Font(size=14, bold=is_bold, italic=is_italic, underline='single' if is_underline else None)
    title_cell.alignment = Alignment(horizontal='center' if is_centered else 'left', vertical='center')
    ws.row_dimensions[1].height = 30
    
    # Empty row 2 for spacing
    ws.row_dimensions[2].height = 10
    
    # Headers row (row 3)
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    
    # Data rows - start from row 4
    for row_idx, product in enumerate(products, 4):
        ws.cell(row=row_idx, column=1, value=row_idx - 3)  # Sequential N° starting from 1
        ws.cell(row=row_idx, column=2, value=product.clean_code or "")
        ws.cell(row=row_idx, column=3, value=product.description or "")
        ws.cell(row=row_idx, column=4, value=product.brand or "")
        
        # Calculate final price - use final_price if set, otherwise use price_usd
        base_price = float(product.price_usd or 0)
        final = float(product.final_price) if product.final_price else base_price
        
        if view_mode == "enterprise":
            ws.cell(row=row_idx, column=5, value=product.original_list_name or "")
            ws.cell(row=row_idx, column=6, value=base_price)
            margin = float(product.margin_percentage) if product.margin_percentage else 0
            ws.cell(row=row_idx, column=7, value=f"{margin:.1f}%")
            ws.cell(row=row_idx, column=8, value=final)  # Always include final price
        else:
            ws.cell(row=row_idx, column=5, value=final)
    
    # Adjust column widths (skip merged cells)
    for col_idx, col in enumerate(ws.columns, 1):
        try:
            # Get column letter from first non-merged cell
            col_letter = openpyxl.utils.get_column_letter(col_idx)
            # Calculate max length, skipping merged cells
            max_length = 0
            for cell in col:
                if cell.value and not isinstance(cell, openpyxl.cell.cell.MergedCell):
                    max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = min(max_length + 2, 50)
        except Exception:
            pass  # Skip on any error
    
    # Save to bytes
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    filename = f"master_table_{view_mode}.xlsx"
    
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


async def _export_to_pdf(products: List[MasterProduct], view_mode: str, custom_title: str = "Listado de Productos"):
    """Generate PDF export."""
    from fastapi.responses import StreamingResponse
    import io
    import re
    
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter, landscape
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    except ImportError:
        raise HTTPException(
            status_code=500, 
            detail="reportlab not installed. Run: pip install reportlab"
        )
    
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter))
    elements = []
    
    styles = getSampleStyleSheet()
    
    # Convert Tiptap HTML to ReportLab-compatible HTML
    # Tiptap uses <strong> for bold, ReportLab uses <b>
    # Tiptap uses <em> for italic, ReportLab uses <i>
    title_html = custom_title
    title_html = title_html.replace('<strong>', '<b>').replace('</strong>', '</b>')
    title_html = title_html.replace('<em>', '<i>').replace('</em>', '</i>')
    # Remove or replace unsupported tags
    title_html = title_html.replace('<br>', ' ').replace('<br/>', ' ').replace('<br />', ' ')
    # Remove paragraph tags but keep content
    title_html = re.sub(r'<p[^>]*>', '', title_html).replace('</p>', ' ')
    
    # Title - pass HTML directly (ReportLab Paragraph supports <b>, <i>, <u>)
    title = Paragraph(
        title_html,
        styles['Title']
    )
    elements.append(title)
    elements.append(Spacer(1, 12))
    
    # Headers - always include final price in enterprise view
    if view_mode == "enterprise":
        headers = ["N°", "CÓDIGO", "DESCRIPCIÓN", "MARCA", "LISTA", "USD", "%", "USD FINAL"]
    else:
        headers = ["N°", "CÓDIGO", "DESCRIPCIÓN", "MARCA", "USD"]
    
    data = [headers]
    
    # Data rows - no limit, use sequential numbering
    for idx, product in enumerate(products, 1):
        base_price = float(product.price_usd or 0)
        final = float(product.final_price) if product.final_price else base_price
        margin = float(product.margin_percentage) if product.margin_percentage else 0
        
        if view_mode == "enterprise":
            row = [
                idx,  # Sequential N°
                (product.clean_code or "")[:20],
                (product.description or "")[:40],
                (product.brand or "")[:15],
                (product.original_list_name or "")[:15],
                f"${base_price:.2f}",
                f"{margin:.1f}%",
                f"${final:.2f}"  # Always include final price
            ]
        else:
            row = [
                idx,  # Sequential N°
                (product.clean_code or "")[:20],
                (product.description or "")[:50],
                (product.brand or "")[:15],
                f"${final:.2f}"
            ]
        data.append(row)
    
    # Create table
    table = Table(data)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.white),
        ('TEXTCOLOR', (0, 1), (-1, -1), colors.black),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
    ]))
    
    elements.append(table)
    doc.build(elements)
    
    buffer.seek(0)
    filename = f"master_table_{view_mode}.pdf"
    
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


async def _export_to_excel_with_header(
    products: List[MasterProduct],
    view_mode: str,
    header_html: str,
    filename: str
):
    """
    Generate Excel export with HTML header supporting per-paragraph alignment.
    Separates filename from header content.
    """
    from fastapi.responses import StreamingResponse
    import re
    
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise HTTPException(
            status_code=500, 
            detail="openpyxl not installed. Run: pip install openpyxl"
        )
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = filename[:31] if filename else "Listado"
    
    # Parse HTML content - split by paragraphs (capture full tag + content)
    # Match pattern: <p style="...">content</p>
    paragraph_matches = re.findall(r'(<p[^>]*>)(.*?)</p>', header_html, re.DOTALL)
    if not paragraph_matches:
        # If no <p> tags, treat entire content as one paragraph
        paragraph_matches = [('', header_html)]
    
    # Define headers based on view mode
    if view_mode == "enterprise":
        headers = ["N°", "CÓDIGO O REFERENCIA", "DESCRIPCIÓN", "MARCA", "LISTA ORIGINAL", "USD", "%", "USD FINAL"]
    else:
        headers = ["N°", "CÓDIGO O REFERENCIA", "DESCRIPCIÓN", "MARCA", "USD"]
    
    current_row = 1
    
    # Process each paragraph - also split by <br> tags for line breaks
    for p_tag, content in paragraph_matches:
        # Extract alignment from the <p> tag itself (not content)
        align_match = re.search(r'text-align:\s*(left|center|right)', p_tag)
        alignment = align_match.group(1) if align_match else 'left'
        
        # Split by <br> tags (handles <br>, <br/>, <br />)
        lines = re.split(r'<br\s*/?>', content)
        
        for line in lines:
            line = line.strip()
            if not line:
                # Empty line (from consecutive <br>) - add spacing row
                ws.row_dimensions[current_row].height = 10
                current_row += 1
                continue
                
            # Parse formatting from this line
            is_bold = '<strong>' in line or '<b>' in line
            is_italic = '<em>' in line or '<i>' in line
            is_underline = '<u>' in line
            
            # Clean text (strip HTML tags)
            text = re.sub(r'<[^>]+>', '', line).strip()
            
            if text:
                # Merge cells for this line
                ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(headers))
                cell = ws.cell(row=current_row, column=1, value=text)
                cell.font = Font(size=12, bold=is_bold, italic=is_italic, underline='single' if is_underline else None)
                cell.alignment = Alignment(horizontal=alignment, vertical='center')
                ws.row_dimensions[current_row].height = 22
                current_row += 1
    
    # Empty row for spacing
    if current_row > 1:
        ws.row_dimensions[current_row].height = 10
        current_row += 1
    
    # Headers row
    header_row = current_row
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    
    # Data rows
    data_start = header_row + 1
    for row_idx, product in enumerate(products, data_start):
        seq_num = row_idx - header_row
        ws.cell(row=row_idx, column=1, value=seq_num)
        ws.cell(row=row_idx, column=2, value=product.clean_code or "")
        ws.cell(row=row_idx, column=3, value=product.description or "")
        ws.cell(row=row_idx, column=4, value=product.brand or "")
        
        base_price = float(product.price_usd or 0)
        final = float(product.final_price) if product.final_price else base_price
        
        if view_mode == "enterprise":
            ws.cell(row=row_idx, column=5, value=product.original_list_name or "")
            ws.cell(row=row_idx, column=6, value=base_price)
            margin = float(product.margin_percentage) if product.margin_percentage else 0
            ws.cell(row=row_idx, column=7, value=f"{margin:.1f}%")
            ws.cell(row=row_idx, column=8, value=final)
        else:
            ws.cell(row=row_idx, column=5, value=final)
    
    # Adjust column widths
    for col_idx, col in enumerate(ws.columns, 1):
        try:
            col_letter = openpyxl.utils.get_column_letter(col_idx)
            max_length = 0
            for cell in col:
                if cell.value and not isinstance(cell, openpyxl.cell.cell.MergedCell):
                    max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = min(max_length + 2, 50)
        except Exception:
            pass
    
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    # Add date to filename
    from datetime import datetime
    date_str = datetime.now().strftime('%Y-%m-%d')
    output_filename = f"{filename.replace(' ', '_')}_{date_str}.xlsx"
    
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={output_filename}"}
    )


async def _export_to_pdf_with_header(
    products: List[MasterProduct],
    view_mode: str,
    header_html: str,
    filename: str
):
    """
    Generate PDF export with HTML header supporting per-paragraph alignment.
    Separates filename from header content.
    """
    from fastapi.responses import StreamingResponse
    import re
    
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter, landscape
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT
    except ImportError:
        raise HTTPException(
            status_code=500, 
            detail="reportlab not installed. Run: pip install reportlab"
        )
    
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter))
    elements = []
    
    styles = getSampleStyleSheet()
    
    # Parse HTML content - split by paragraphs (capture full tag + content)
    paragraph_matches = re.findall(r'(<p[^>]*>)(.*?)</p>', header_html, re.DOTALL)
    if not paragraph_matches:
        paragraph_matches = [('', header_html)]
    
    # Process each paragraph - also split by <br> tags for line breaks
    for p_tag, content in paragraph_matches:
        # Extract alignment from the <p> tag itself
        align_match = re.search(r'text-align:\s*(left|center|right)', p_tag)
        alignment = align_match.group(1) if align_match else 'left'
        
        # Split by <br> tags (handles <br>, <br/>, <br />)
        lines = re.split(r'<br\s*/?>', content)
        
        for line in lines:
            line = line.strip()
            if not line:
                # Empty line - add spacer
                elements.append(Spacer(1, 8))
                continue
            
            # Convert HTML tags
            para_html = line
            para_html = para_html.replace('<strong>', '<b>').replace('</strong>', '</b>')
            para_html = para_html.replace('<em>', '<i>').replace('</em>', '</i>')
            
            text = para_html.strip()
            
            if text:
                # Create style with proper alignment
                align_map = {'left': TA_LEFT, 'center': TA_CENTER, 'right': TA_RIGHT}
                para_style = ParagraphStyle(
                    'CustomPara',
                    parent=styles['Title'],
                    alignment=align_map.get(alignment, TA_LEFT),
                    fontSize=14,
                    spaceAfter=6
                )
                elements.append(Paragraph(text, para_style))
    
    elements.append(Spacer(1, 12))
    
    # Headers
    if view_mode == "enterprise":
        headers = ["N°", "CÓDIGO", "DESCRIPCIÓN", "MARCA", "LISTA", "USD", "%", "USD FINAL"]
    else:
        headers = ["N°", "CÓDIGO", "DESCRIPCIÓN", "MARCA", "USD"]
    
    data = [headers]
    
    # Data rows
    for idx, product in enumerate(products, 1):
        base_price = float(product.price_usd or 0)
        final = float(product.final_price) if product.final_price else base_price
        margin = float(product.margin_percentage) if product.margin_percentage else 0
        
        if view_mode == "enterprise":
            row = [
                idx,
                (product.clean_code or "")[:20],
                (product.description or "")[:40],
                (product.brand or "")[:15],
                (product.original_list_name or "")[:15],
                f"${base_price:.2f}",
                f"{margin:.1f}%",
                f"${final:.2f}"
            ]
        else:
            row = [
                idx,
                (product.clean_code or "")[:20],
                (product.description or "")[:50],
                (product.brand or "")[:15],
                f"${final:.2f}"
            ]
        data.append(row)
    
    table = Table(data)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.white),
        ('TEXTCOLOR', (0, 1), (-1, -1), colors.black),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
    ]))
    
    elements.append(table)
    doc.build(elements)
    
    buffer.seek(0)
    
    # Add date to filename
    from datetime import datetime
    date_str = datetime.now().strftime('%Y-%m-%d')
    output_filename = f"{filename.replace(' ', '_')}_{date_str}.pdf"
    
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={output_filename}"}
    )


async def _export_to_excel_with_template(
    products: List[MasterProduct],
    view_mode: str,
    template: TemplateFileConfig,
    filename: str
):
    """
    Generate Excel export using uploaded template as header.
    
    Loads the template, finds the last row with data, and inserts
    product data below it.
    """
    from fastapi.responses import StreamingResponse
    
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise HTTPException(
            status_code=500, 
            detail="openpyxl not installed. Run: pip install openpyxl"
        )
    
    # Decode template from base64
    try:
        template_bytes = base64.b64decode(template.data)
        template_buffer = io.BytesIO(template_bytes)
        wb = openpyxl.load_workbook(template_buffer)
        ws = wb.active
    except Exception as e:
        logger.error(f"Failed to load Excel template: {e}")
        raise HTTPException(status_code=400, detail=f"Invalid Excel template: {str(e)}")
    
    # Find the last row with data
    last_row = ws.max_row
    
    # Add spacing row
    start_row = last_row + 2
    
    # Headers based on view mode
    if view_mode == "enterprise":
        headers = ["N°", "CÓDIGO O REFERENCIA", "DESCRIPCIÓN", "MARCA", "LISTA ORIGINAL", "USD", "%", "USD FINAL"]
    else:
        headers = ["N°", "CÓDIGO O REFERENCIA", "DESCRIPCIÓN", "MARCA", "USD"]
    
    # Style header
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    
    # Data rows
    for row_idx, product in enumerate(products, start_row + 1):
        seq_num = row_idx - start_row
        ws.cell(row=row_idx, column=1, value=seq_num)
        ws.cell(row=row_idx, column=2, value=product.clean_code or "")
        ws.cell(row=row_idx, column=3, value=product.description or "")
        ws.cell(row=row_idx, column=4, value=product.brand or "")
        
        base_price = float(product.price_usd or 0)
        final = float(product.final_price) if product.final_price else base_price
        
        if view_mode == "enterprise":
            ws.cell(row=row_idx, column=5, value=product.original_list_name or "")
            ws.cell(row=row_idx, column=6, value=base_price)
            margin = float(product.margin_percentage) if product.margin_percentage else 0
            ws.cell(row=row_idx, column=7, value=f"{margin:.1f}%")
            ws.cell(row=row_idx, column=8, value=final)
        else:
            ws.cell(row=row_idx, column=5, value=final)
    
    # Adjust column widths for new data
    for col in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 50)
    
    # Save to bytes
    output_buffer = io.BytesIO()
    wb.save(output_buffer)
    output_buffer.seek(0)
    
    output_filename = f"{filename.replace(' ', '_')}.xlsx"
    
    return StreamingResponse(
        output_buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={output_filename}"}
    )


async def _export_to_pdf_with_template(
    products: List[MasterProduct],
    view_mode: str,
    template: TemplateFileConfig,
    filename: str
):
    """
    Generate PDF export using uploaded template as header pages.
    
    Loads the template PDF and appends the data table as additional pages.
    """
    from fastapi.responses import StreamingResponse
    
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter, landscape
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Spacer
        from pypdf import PdfReader, PdfWriter
    except ImportError as e:
        raise HTTPException(
            status_code=500, 
            detail=f"Required library not installed: {e}. Run: pip install reportlab pypdf"
        )
    
    # First, generate the data table PDF
    data_buffer = io.BytesIO()
    doc = SimpleDocTemplate(data_buffer, pagesize=landscape(letter))
    elements = []
    
    # Add some spacing at top
    elements.append(Spacer(1, 20))
    
    # Headers
    if view_mode == "enterprise":
        headers = ["N°", "CÓDIGO", "DESCRIPCIÓN", "MARCA", "LISTA", "USD", "%", "USD FINAL"]
    else:
        headers = ["N°", "CÓDIGO", "DESCRIPCIÓN", "MARCA", "USD"]
    
    data = [headers]
    
    # Data rows
    for idx, product in enumerate(products, 1):
        base_price = float(product.price_usd or 0)
        final = float(product.final_price) if product.final_price else base_price
        margin = float(product.margin_percentage) if product.margin_percentage else 0
        
        if view_mode == "enterprise":
            row = [
                idx,
                (product.clean_code or "")[:20],
                (product.description or "")[:40],
                (product.brand or "")[:15],
                (product.original_list_name or "")[:15],
                f"${base_price:.2f}",
                f"{margin:.1f}%",
                f"${final:.2f}"
            ]
        else:
            row = [
                idx,
                (product.clean_code or "")[:20],
                (product.description or "")[:50],
                (product.brand or "")[:15],
                f"${final:.2f}"
            ]
        data.append(row)
    
    # Create table
    table = Table(data)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.white),
        ('TEXTCOLOR', (0, 1), (-1, -1), colors.black),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
    ]))
    
    elements.append(table)
    doc.build(elements)
    data_buffer.seek(0)
    
    # Now merge with template
    try:
        # Load template PDF
        template_bytes = base64.b64decode(template.data)
        template_reader = PdfReader(io.BytesIO(template_bytes))
        
        # Load data PDF
        data_reader = PdfReader(data_buffer)
        
        # Create output
        writer = PdfWriter()
        
        # Add template pages first
        for page in template_reader.pages:
            writer.add_page(page)
        
        # Add data pages
        for page in data_reader.pages:
            writer.add_page(page)
        
        # Write to buffer
        output_buffer = io.BytesIO()
        writer.write(output_buffer)
        output_buffer.seek(0)
        
    except Exception as e:
        logger.error(f"Failed to merge PDF template: {e}")
        # Fall back to just the data PDF
        data_buffer.seek(0)
        output_buffer = data_buffer
    
    output_filename = f"{filename.replace(' ', '_')}.pdf"
    
    return StreamingResponse(
        output_buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={output_filename}"}
    )

